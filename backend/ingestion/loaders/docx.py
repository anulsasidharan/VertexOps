"""DOCX document loader — requires python-docx (pip install python-docx)."""

import io
from typing import Optional

from backend.ingestion.loaders.base import BaseLoader, ParsedDocument

_DOCX_AVAILABLE: Optional[bool] = None


def _require_docx():
    global _DOCX_AVAILABLE
    try:
        import docx  # noqa: F401

        _DOCX_AVAILABLE = True
        return docx
    except ImportError:
        _DOCX_AVAILABLE = False
        raise ImportError(
            "python-docx is required for DOCX parsing. Install it with: pip install python-docx"
        )


class DocxLoader(BaseLoader):
    def load(self, content: bytes, filename: str) -> ParsedDocument:
        docx = _require_docx()

        doc = docx.Document(io.BytesIO(content))

        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        text = "\n\n".join(paragraphs)

        # Use the first paragraph as a title candidate (often the heading)
        title: Optional[str] = paragraphs[0][:200] if paragraphs else None

        # Check core properties for a real title
        try:
            core_title = doc.core_properties.title
            if core_title and core_title.strip():
                title = core_title.strip()[:200]
        except Exception:
            pass

        return ParsedDocument(text=text, format="docx", title=title)
